"""Excel and CSV exporters for financial document records."""

import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from hesabyar.core.exceptions import ExportError
from hesabyar.core.logger import get_logger
from hesabyar.models.schemas import ExtractedDocument, FileRecord, FileStatus

logger = get_logger("exporters.excel_csv")


def records_to_rows(records: List[FileRecord]) -> List[Dict[str, Any]]:
    """Transform a list of FileRecord instances into flat tabular dictionaries."""
    rows: List[Dict[str, Any]] = []

    for r in records:
        doc: Optional[ExtractedDocument] = None
        if r.extracted_data_json:
            try:
                doc = ExtractedDocument.model_validate_json(r.extracted_data_json)
            except Exception:
                pass

        status_fa = {
            FileStatus.COMPLETED: "معتبر ✓",
            FileStatus.NEEDS_REVIEW: "نیازمند بازبینی ⚠️",
            FileStatus.FAILED: "خطادار ✗",
            FileStatus.PENDING: "در انتظار صف",
            FileStatus.PREPROCESSING: "پیش‌پردازش",
            FileStatus.OCR_PROCESSING: "استخراج OCR",
            FileStatus.EXTRACTING: "استخراج فیلدها",
        }.get(r.status, str(r.status.value))

        currency_fa = "-"
        if doc and doc.currency:
            currency_fa = "تومان" if doc.currency.value == "IRT" else ("ریال" if doc.currency.value == "IRR" else "-")

        rows.append({
            "نام فایل": r.file_name,
            "وضعیت": status_fa,
            "نوع سند": doc.document_type if doc else "-",
            "فروشگاه / پذیرنده": doc.vendor_name if doc and doc.vendor_name else "-",
            "تاریخ سند": doc.issue_date if doc and doc.issue_date else "-",
            "ساعت": doc.issue_time if doc and doc.issue_time else "-",
            "مبلغ کل": doc.total_amount if doc and doc.total_amount is not None else None,
            "واحد پول": currency_fa,
            "کد پیگیری": doc.tracking_code if doc and doc.tracking_code else "-",
            "شماره کارت": doc.card_number if doc and doc.card_number else "-",
            "شماره پایانه": doc.terminal_id if doc and doc.terminal_id else "-",
            "دقت OCR": f"{r.ocr_confidence:.1%}" if r.ocr_confidence is not None else "-",
            "دلایل بازبینی / خطا": r.error_message if r.error_message else "-",
            "تاریخ پردازش": r.updated_at.strftime("%Y-%m-%d %H:%M") if r.updated_at else "-",
        })

    return rows


def export_to_excel(
    records: List[FileRecord], output_path: Union[str, Path]
) -> Path:
    """
    Export document records to an RTL formatted, styled Microsoft Excel spreadsheet.
    """
    path = Path(output_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    rows = records_to_rows(records)
    df = pd.DataFrame(rows)

    try:
        with pd.ExcelWriter(str(path), engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="گزارش اسناد مالی", index=False)
            ws = writer.sheets["گزارش اسناد مالی"]

            # Enable Right-to-Left (RTL) mode for Persian layout
            ws.views.sheetView[0].rightToLeft = True

            # Style Header Row
            header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
            header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for col_idx in range(1, len(df.columns) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align

            # Style Data Rows and Format Numbers
            amount_col_idx = list(df.columns).index("مبلغ کل") + 1 if "مبلغ کل" in df.columns else None

            body_font = Font(name="Tahoma", size=10)
            center_align = Alignment(horizontal="center", vertical="center")
            right_align = Alignment(horizontal="right", vertical="center")

            for row_idx in range(2, len(rows) + 2):
                for col_idx in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = body_font
                    if col_idx == amount_col_idx:
                        cell.number_format = "#,##0"
                        cell.alignment = right_align
                    else:
                        cell.alignment = center_align

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        logger.info(f"Excel export saved: {path}")
        return path
    except Exception as e:
        logger.error(f"Failed exporting to Excel: {e}")
        raise ExportError(f"Excel export error: {e}") from e


def export_to_csv(
    records: List[FileRecord], output_path: Union[str, Path]
) -> Path:
    """
    Export document records to CSV with UTF-8 BOM encoding for seamless Excel compatibility.
    """
    path = Path(output_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    rows = records_to_rows(records)
    df = pd.DataFrame(rows)

    try:
        # utf-8-sig ensures Windows Excel opens Persian characters cleanly without mojibake
        df.to_csv(str(path), index=False, encoding="utf-8-sig")
        logger.info(f"CSV export saved: {path}")
        return path
    except Exception as e:
        logger.error(f"Failed exporting to CSV: {e}")
        raise ExportError(f"CSV export error: {e}") from e
