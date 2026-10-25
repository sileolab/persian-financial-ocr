"""Microsoft Word (DOCX) executive report generator for financial document processing."""

from datetime import datetime
from pathlib import Path
from typing import List, Union

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from hesabyar.core.exceptions import ExportError
from hesabyar.core.logger import get_logger
from hesabyar.models.schemas import ExtractedDocument, FileRecord, FileStatus

logger = get_logger("exporters.word")


def set_cell_background(cell, hex_color: str):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def export_to_word(
    records: List[FileRecord], output_path: Union[str, Path]
) -> Path:
    """
    Generate an executive Word (DOCX) summary report for processed financial receipts.
    """
    path = Path(output_path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    try:
        doc = Document()

        # Page margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_title = p_title.add_run("گـزارش جامع پـردازش اسـناد و رسـیدهای مـالی")
        r_title.font.name = "Tahoma"
        r_title.font.size = Pt(18)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

        # Subtitle
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        r_sub = p_sub.add_run(f"سامانه هوشمند حسابیار (Hesabyar) | تاریخ صدور گزارش: {now_str}")
        r_sub.font.name = "Tahoma"
        r_sub.font.size = Pt(10)
        r_sub.font.italic = True
        r_sub.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

        # Calculate statistics
        total = len(records)
        completed = sum(1 for r in records if r.status == FileStatus.COMPLETED)
        needs_review = sum(1 for r in records if r.status == FileStatus.NEEDS_REVIEW)
        failed = sum(1 for r in records if r.status == FileStatus.FAILED)

        total_amount_sum = 0
        for r in records:
            if r.status == FileStatus.COMPLETED and r.extracted_data_json:
                try:
                    edoc = ExtractedDocument.model_validate_json(r.extracted_data_json)
                    if edoc.total_amount:
                        total_amount_sum += edoc.total_amount
                except Exception:
                    pass

        # Section 1: Summary KPI Table
        p_sec1 = doc.add_paragraph()
        p_sec1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_sec1 = p_sec1.add_run("۱. خـلاصـه آمـاری پـردازش")
        r_sec1.font.name = "Tahoma"
        r_sec1.font.size = Pt(13)
        r_sec1.font.bold = True

        stat_table = doc.add_table(rows=2, cols=5)
        stat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers_stat = ["مجموع مبالغ تاییدشده", "خطادار ✗", "نیازمند بازبینی ⚠️", "تاییدشده ✓", "کل اسناد"]
        vals_stat = [f"{total_amount_sum:,}", str(failed), str(needs_review), str(completed), str(total)]

        for i, (h, v) in enumerate(zip(headers_stat, vals_stat)):
            hdr_cell = stat_table.cell(0, i)
            hdr_cell.text = h
            hdr_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_background(hdr_cell, "1F497D")
            hdr_cell.paragraphs[0].runs[0].font.name = "Tahoma"
            hdr_cell.paragraphs[0].runs[0].font.bold = True
            hdr_cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            val_cell = stat_table.cell(1, i)
            val_cell.text = v
            val_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_background(val_cell, "F2F4F7")
            val_cell.paragraphs[0].runs[0].font.name = "Tahoma"
            val_cell.paragraphs[0].runs[0].font.bold = True
            val_cell.paragraphs[0].runs[0].font.size = Pt(11)

        doc.add_paragraph()

        # Section 2: Detail Table
        p_sec2 = doc.add_paragraph()
        p_sec2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_sec2 = p_sec2.add_run("۲. جـدول جـزئیات اسـناد پـردازش‌شده")
        r_sec2.font.name = "Tahoma"
        r_sec2.font.size = Pt(13)
        r_sec2.font.bold = True

        detail_table = doc.add_table(rows=1, cols=7)
        detail_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        col_titles = ["وضعیت", "کد پیگیری", "مبلغ", "تاریخ", "فروشگاه", "نام فایل", "ردیف"]
        for j, title in enumerate(col_titles):
            cell = detail_table.cell(0, j)
            cell.text = title
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_background(cell, "2B579A")
            run = cell.paragraphs[0].runs[0]
            run.font.name = "Tahoma"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for idx, rec in enumerate(records, start=1):
            row_cells = detail_table.add_row().cells
            doc_data: Optional[ExtractedDocument] = None
            if rec.extracted_data_json:
                try:
                    doc_data = ExtractedDocument.model_validate_json(rec.extracted_data_json)
                except Exception:
                    pass

            vendor = doc_data.vendor_name if doc_data and doc_data.vendor_name else "-"
            date_val = doc_data.issue_date if doc_data and doc_data.issue_date else "-"
            amount_val = f"{doc_data.total_amount:,}" if doc_data and doc_data.total_amount is not None else "-"
            tracking = doc_data.tracking_code if doc_data and doc_data.tracking_code else "-"
            status_text = "معتبر ✓" if rec.status == FileStatus.COMPLETED else (
                "بازبینی ⚠️" if rec.status == FileStatus.NEEDS_REVIEW else "خطا ✗"
            )

            row_data = [status_text, tracking, amount_val, date_val, vendor, rec.file_name, str(idx)]

            for col_idx, text_val in enumerate(row_data):
                c = row_cells[col_idx]
                c.text = text_val
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_elem = c.paragraphs[0].runs[0]
                r_elem.font.name = "Tahoma"
                r_elem.font.size = Pt(9)
                if idx % 2 == 0:
                    set_cell_background(c, "F9FAFB")

        # Section 3: Review Notes if any
        review_records = [r for r in records if r.status == FileStatus.NEEDS_REVIEW]
        if review_records:
            doc.add_paragraph()
            p_sec3 = doc.add_paragraph()
            p_sec3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_sec3 = p_sec3.add_run("۳. اسـناد نیازمند بازبینی دستی (Discrepancies & Review Reasons)")
            r_sec3.font.name = "Tahoma"
            r_sec3.font.size = Pt(13)
            r_sec3.font.bold = True
            r_sec3.font.color.rgb = RGBColor(0xB2, 0x59, 0x00)

            for rr in review_records:
                p_item = doc.add_paragraph()
                p_item.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r_bold = p_item.add_run(f"• سند: {rr.file_name} -> ")
                r_bold.font.name = "Tahoma"
                r_bold.font.bold = True
                r_desc = p_item.add_run(rr.error_message or "عدم انطباق محاسباتی یا کسری اطلاعات")
                r_desc.font.name = "Tahoma"
                r_desc.font.size = Pt(9.5)

        doc.save(str(path))
        logger.info(f"Word executive report saved: {path}")
        return path
    except Exception as e:
        logger.error(f"Failed exporting to Word: {e}")
        raise ExportError(f"Word export error: {e}") from e
