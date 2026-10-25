"""Unit tests for reporting and export modules (Excel, Word, CSV, JSON)."""

import json
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd
from docx import Document
import pytest

from hesabyar.exporters import (
    export_to_csv,
    export_to_excel,
    export_to_json,
    export_to_word,
    generate_summary_report,
)
from hesabyar.models.schemas import CurrencyUnit, ExtractedDocument, FileRecord, FileStatus


@pytest.fixture
def sample_records():
    """Create a list of mock FileRecord items with financial data."""
    doc1 = ExtractedDocument(
        document_type="pos_slip",
        vendor_name="هایپرمارکت بهاران",
        issue_date="1403/08/04",
        issue_time="18:30:00",
        total_amount=250000,
        currency=CurrencyUnit.IRT,
        tracking_code="123456",
        card_number="6037-9971-****-1234",
    )

    doc2 = ExtractedDocument(
        document_type="invoice",
        vendor_name="شرکت رایان",
        issue_date="1403/08/01",
        total_amount=5000000,
        currency=CurrencyUnit.IRR,
        needs_review=True,
        review_reasons=["Mismatched line item totals"],
    )

    rec1 = FileRecord(
        file_hash="hash11111111",
        file_path="data/input/receipt_01.jpg",
        file_name="receipt_01.jpg",
        file_size=1024,
        status=FileStatus.COMPLETED,
        ocr_confidence=0.96,
        extracted_data_json=doc1.model_dump_json(),
        created_at=datetime.now(),
        updated_at=datetime.now(),
    )

    rec2 = FileRecord(
        file_hash="hash22222222",
        file_path="data/input/invoice_02.png",
        file_name="invoice_02.png",
        file_size=2048,
        status=FileStatus.NEEDS_REVIEW,
        ocr_confidence=0.88,
        extracted_data_json=doc2.model_dump_json(),
        error_message="Mismatched line item totals",
        created_at=datetime.now(),
        updated_at=datetime.now(),
    )

    return [rec1, rec2]


def test_export_to_excel(tmp_path: Path, sample_records):
    excel_path = tmp_path / "test_report.xlsx"
    out = export_to_excel(sample_records, excel_path)
    assert out.exists()

    # Open with openpyxl and inspect
    wb = openpyxl.load_workbook(str(out))
    ws = wb.active
    assert ws.views.sheetView[0].rightToLeft is True
    assert ws.cell(row=1, column=1).value == "نام فایل"
    assert ws.cell(row=2, column=1).value == "receipt_01.jpg"
    assert ws.cell(row=2, column=7).value == 250000  # Amount column


def test_export_to_csv(tmp_path: Path, sample_records):
    csv_path = tmp_path / "test_report.csv"
    out = export_to_csv(sample_records, csv_path)
    assert out.exists()

    # Verify UTF-8 BOM is present
    raw_bytes = csv_path.read_bytes()
    assert raw_bytes.startswith(b"\xef\xbb\xbf")

    # Read back with Pandas
    df = pd.read_csv(str(csv_path), encoding="utf-8-sig")
    assert len(df) == 2
    assert "نام فایل" in df.columns
    assert "هایپرمارکت بهاران" in df["فروشگاه / پذیرنده"].values


def test_export_to_word(tmp_path: Path, sample_records):
    word_path = tmp_path / "test_report.docx"
    out = export_to_word(sample_records, word_path)
    assert out.exists()

    doc = Document(str(out))
    assert len(doc.tables) >= 2
    # Verify title text
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "گزارش جامع" in full_text or "گـزارش جامع" in full_text
    assert "هایپرمارکت بهاران" in " ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)


def test_generate_summary_report_and_json(tmp_path: Path, sample_records):
    json_path = tmp_path / "test_report.json"
    out = export_to_json(sample_records, json_path)
    assert out.exists()

    data = json.loads(json_path.read_text(encoding="utf-8"))
    assert "kpis" in data
    assert data["kpis"]["total_files"] == 2
    assert data["kpis"]["completed_verified"] == 1
    assert data["kpis"]["needs_review"] == 1

    # Financial totals
    fin = data["financial_totals"]
    assert fin["total_toman_irt"] == 250000
    # Equivalent Toman: 250,000 IRT + (0 IRR from completed since doc2 is in review)
    assert fin["equivalent_total_toman"] == 250000
